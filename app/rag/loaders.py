from __future__ import annotations

import csv
import json
import re
import unicodedata
import pymupdf
import zipfile
from xml.etree import ElementTree
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook

from .models import Document


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".jsonl",
    ".docx",
    ".odt",
    ".rtf",
    ".xlsx",
    ".html",
    ".htm",
}


def discover_files(data_dir: Path) -> list[Path]:
    return sorted(
        path
        for path in data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def load_documents(path: Path, root: Path) -> list[Document]:
    documents = _load_documents(path, root)
    if len(documents) > 2000 or sum(len(d.text) for d in documents) > 6000000:
        raise ValueError("Document extraction exceeds the 2,000-page or 6-million-character limit.")
    return documents


def _load_documents(path: Path, root: Path) -> list[Document]:
    ext = path.suffix.lower()
    rel_path = str(path.relative_to(root))

    if ext == ".pdf":
        return list(_load_pdf(path, rel_path))
    if ext in {".txt", ".md", ".markdown"}:
        return [_plain_document(path, rel_path)]
    if ext == ".csv":
        return [_load_csv(path, rel_path)]
    if ext == ".json":
        return [_load_json(path, rel_path)]
    if ext == ".jsonl":
        return [_load_jsonl(path, rel_path)]
    if ext == ".docx":
        return [_load_docx(path, rel_path)]
    if ext == ".odt":
        return [_load_odt(path, rel_path)]
    if ext == ".rtf":
        return [_load_rtf(path, rel_path)]
    if ext == ".xlsx":
        return [_load_xlsx(path, rel_path)]
    if ext in {".html", ".htm"}:
        return [_load_html(path, rel_path)]

    return []


def _plain_document(path: Path, rel_path: str) -> Document:
    return Document(
        text=path.read_text(encoding="utf-8-sig", errors="ignore"),
        source=path.name,
        path=rel_path,
    )


def _load_pdf(path: Path, rel_path: str) -> Iterable[Document]:
    # MuPDF preserves embedded-font Unicode mappings and combining marks that
    # the previous parser sometimes replaced with punctuation or extra spaces.
    with pymupdf.open(stream=path.read_bytes(), filetype="pdf") as pdf:
        if pdf.needs_pass:
            raise ValueError('This PDF is encrypted. Upload an unlocked copy.')
        if len(pdf) > 2000:
            raise ValueError("PDF exceeds the 2,000-page limit.")
        extracted = 0
        for number, page in enumerate(pdf, start=1):
            text = unicodedata.normalize('NFC', page.get_text('text', sort=False))
            extracted += len(text)
            if extracted > 6000000:
                raise ValueError('PDF text exceeds the extraction limit.')
            if text.strip():
                yield Document(text=text, source=path.name, path=rel_path, page=number)


def _load_csv(path: Path, rel_path: str) -> Document:
    rows: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
        reader = csv.DictReader(handle, dialect=dialect)
        if reader.fieldnames:
            for row_number, row in enumerate(reader, start=1):
                fields = [f"{key}: {value}" for key, value in row.items() if value]
                if fields:
                    rows.append(f"Row {row_number}: " + "; ".join(fields))
        else:
            handle.seek(0)
            for row_number, row in enumerate(csv.reader(handle, dialect=dialect), start=1):
                rows.append(f"Row {row_number}: " + "; ".join(row))
    return Document(text="\n".join(rows), source=path.name, path=rel_path)


def _load_json(path: Path, rel_path: str) -> Document:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return Document(text=json.dumps(data, ensure_ascii=True, indent=2), source=path.name, path=rel_path)


def _load_jsonl(path: Path, rel_path: str) -> Document:
    lines: list[str] = []
    for index, line in enumerate(path.read_text(encoding="utf-8", errors="ignore").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            parsed = json.loads(line)
            lines.append(f"Record {index}: {json.dumps(parsed, ensure_ascii=True)}")
        except json.JSONDecodeError:
            lines.append(f"Record {index}: {line}")
    return Document(text="\n".join(lines), source=path.name, path=rel_path)


def _load_docx(path: Path, rel_path: str) -> Document:
    doc = DocxDocument(str(path))
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"Table {table_index}:")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return Document(text="\n".join(parts), source=path.name, path=rel_path)


def _load_odt(path: Path, rel_path: str) -> Document:
    with zipfile.ZipFile(path) as archive:
        xml_text = archive.read("content.xml")
    root = ElementTree.fromstring(xml_text)
    parts = [text.strip() for text in root.itertext() if text and text.strip()]
    return Document(text="\n".join(parts), source=path.name, path=rel_path)


def _load_rtf(path: Path, rel_path: str) -> Document:
    text = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ").replace("\\", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return Document(text=text, source=path.name, path=rel_path)


def _load_xlsx(path: Path, rel_path: str) -> Document:
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts, cells = [], 0
    try:
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > 500000:
                    raise ValueError('Spreadsheet exceeds the 500,000-cell extraction limit.')
                values = [str(value) for value in row if value is not None]
                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()
    return Document(text="\n".join(parts), source=path.name, path=rel_path)


def _load_html(path: Path, rel_path: str) -> Document:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return Document(text=soup.get_text("\n", strip=True), source=path.name, path=rel_path)
